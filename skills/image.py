"""Skill de generacion de imagenes: Agnes AI (hq) + Cloudflare Flux (fast)."""
import os
import re
import time
import base64
import io
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from core.model_config import get_model

import ollama
import requests
from dotenv import load_dotenv
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from skills.base import Skill
from core import confirmation
from core.config_loader import CONFIG

ROOT = Path(__file__).resolve().parent.parent
IMAGES_DIR = ROOT / "sandbox" / "images"

load_dotenv(ROOT / ".env")
AGNES_API_KEY = os.getenv("AGNES_API_KEY")
CLOUDFLARE_ACCOUNT_ID = os.getenv("CLOUDFLARE_ACCOUNT_ID")
CLOUDFLARE_API_TOKEN = os.getenv("CLOUDFLARE_API_TOKEN")

DEFAULT_WIDTH = 1024
DEFAULT_HEIGHT = 1024
TIMEOUT = 120


class ImageSkill(Skill):
    name = "image"
    description = "Genera imagenes desde texto (Agnes AI + Cloudflare Flux)"

    def run(self, action, params):
        if action == "generate":
            return self._generate(
                params.get("prompt", ""),
                params.get("width", DEFAULT_WIDTH),
                params.get("height", DEFAULT_HEIGHT),
                params.get("quality", "fast"),
            )
        if action == "to_word":
            return self._to_word(
                params.get("prompt", ""),
                params.get("count", 1),
                params.get("title", ""),
            )
        if action == "list":
            return self._list()
        return f"Accion desconocida en image: {action}"

    # ─── HELPERS ─────────────────────────────────────────────────────────

    def _slug(self, text, maxlen=50):
        s = text.lower()
        for k, v in {"á":"a","é":"e","í":"i","ó":"o","ú":"u","ñ":"n","ü":"u"}.items():
            s = s.replace(k, v)
        s = re.sub(r'[^a-z0-9]+', '_', s)[:maxlen].strip("_")
        return s or "imagen"

    def _save_image(self, content, prompt):
        IMAGES_DIR.mkdir(parents=True, exist_ok=True)
        slug = self._slug(prompt)
        ts = int(datetime.now().timestamp())
        path = IMAGES_DIR / f"{slug}_{ts}.jpg"
        path.write_bytes(content)
        return path

    # ─── MOTOR 1: AGNES AI (alta calidad) ────────────────────────────────

    def _generate_agnes(self, prompt, width, height):
        """Genera con Agnes AI (motor principal de alta calidad)."""
        if not AGNES_API_KEY:
            raise RuntimeError("Falta AGNES_API_KEY en .env")

        # Determinar tier de tamaño
        if width >= 3072 or height >= 3072:
            size = "4K"
        elif width >= 2048 or height >= 2048:
            size = "3K"
        elif width >= 1536 or height >= 1536:
            size = "2K"
        else:
            size = "1K"

        # Ratio basado en las dimensiones
        ratio = "1:1"
        if width > height * 1.2:
            ratio = "16:9"
        elif height > width * 1.2:
            ratio = "9:16"

        payload = {
            "model": "agnes-image-2.5-flash",
            "prompt": prompt,
            "size": size,
            "ratio": ratio,
            "extra_body": {"response_format": "url"},
        }

        headers = {
            "Authorization": f"Bearer {AGNES_API_KEY}",
            "Content-Type": "application/json",
        }

        r = requests.post(
            "https://apihub.agnes-ai.com/v1/images/generations",
            headers=headers,
            json=payload,
            timeout=TIMEOUT,
        )
        if r.status_code != 200:
            raise RuntimeError(f"Agnes AI error {r.status_code}: {r.text[:300]}")

        data = r.json()
        # Extraer URL de la imagen
        img_url = None
        for item in data.get("data", []):
            if item.get("url"):
                img_url = item["url"]
                break
        if not img_url:
            raise RuntimeError(f"Agnes AI no devolvio URL. Respuesta: {str(data)[:300]}")

        # Descargar la imagen
        img_resp = requests.get(img_url, timeout=60)
        if img_resp.status_code != 200:
            raise RuntimeError(f"No pude descargar la imagen de Agnes: {img_resp.status_code}")

        return img_resp.content, img_url

    # ─── MOTOR 2: CLOUDFLARE FLUX (rapido) ───────────────────────────────

    def _generate_cloudflare(self, prompt):
        """Genera con Cloudflare Workers AI (Flux Schnell, rapido)."""
        if not CLOUDFLARE_ACCOUNT_ID or not CLOUDFLARE_API_TOKEN:
            raise RuntimeError("Faltan CLOUDFLARE_ACCOUNT_ID o CLOUDFLARE_API_TOKEN en .env")

        url = (
            f"https://api.cloudflare.com/client/v4/accounts/"
            f"{CLOUDFLARE_ACCOUNT_ID}/ai/run/@cf/black-forest-labs/flux-1-schnell"
        )
        headers = {"Authorization": f"Bearer {CLOUDFLARE_API_TOKEN}"}
        payload = {"prompt": prompt, "steps": 8}

        r = requests.post(url, headers=headers, json=payload, timeout=TIMEOUT)
        if r.status_code != 200:
            raise RuntimeError(f"Cloudflare error {r.status_code}: {r.text[:300]}")

        data = r.json()
        # Cloudflare devuelve base64 en result.image
        b64 = (data.get("result") or {}).get("image")
        if not b64:
            raise RuntimeError(f"Cloudflare no devolvio imagen. Respuesta: {str(data)[:300]}")

        content = base64.b64decode(b64)
        return content, None

    # ─── GENERATE (con enrutamiento por calidad) ─────────────────────────

    def _generate(self, prompt, width, height, quality="fast"):
        prompt = (prompt or "").strip()
        if not prompt:
            return "Dime que imagen quieres que genere."

        try:
            width = int(width)
            height = int(height)
        except (ValueError, TypeError):
            width, height = DEFAULT_WIDTH, DEFAULT_HEIGHT
        width = max(256, min(width, 4096))
        height = max(256, min(height, 4096))

        if quality not in ("fast", "hq"):
            quality = "fast"

        summary = f"Generar imagen ({quality}): {prompt[:80]} ({width}x{height})"
        if not confirmation.require("image", "generate", summary):
            return "Cancelado."

        print(f"[IMAGE] Generando ({quality}) con '{prompt[:60]}'...")

        try:
            t0 = time.time()
            if quality == "hq":
                content, src_url = self._generate_agnes(prompt, width, height)
            else:
                content, src_url = self._generate_cloudflare(prompt)
            elapsed = time.time() - t0
        except Exception as e:
            # Fallback: si hq falla, intentar fast
            if quality == "hq":
                print(f"[IMAGE] Agnes fallo ({e}). Fallback a Cloudflare...")
                try:
                    content, src_url = self._generate_cloudflare(prompt)
                    elapsed = time.time() - t0
                except Exception as e2:
                    return f"Error generando imagen: {e} / fallback: {e2}"
            else:
                return f"Error generando imagen: {e}"

        path = self._save_image(content, prompt)
        size_kb = len(content) // 1024
        print(f"[IMAGE] Guardada: {path} ({size_kb} KB en {elapsed:.1f}s)")

        return {
            "thought": f"Imagen generada ({quality}) en {elapsed:.1f}s",
            "display": (
                f"Imagen generada ({quality.upper()}):\n"
                f"  {path}\n"
                f"  {width}x{height}, {size_kb} KB, {elapsed:.1f}s\n\n"
                f"Prompt: {prompt}"
            ),
            "voice": f"Listo. Imagen guardada en {path.name}.",
        }

    # ─── COMBO: IMAGEN → WORD ────────────────────────────────────────────

    def _to_word(self, prompt, count, title):
        try:
            count = int(count)
        except (ValueError, TypeError):
            count = 1
        count = max(1, min(count, 5))

        imagenes = []

        if prompt:
            print(f"[IMAGE] Generando {count} imagen(es) HQ para '{prompt[:60]}'...")
            for i in range(count):
                variante = prompt if count == 1 else f"{prompt} (variante {i+1})"
                result = self._generate(variante, 1024, 1024, quality="hq")
                if isinstance(result, dict):
                    display = result.get("display", "")
                    m = re.search(r'([A-Za-z]:\\[^\s]+\.jpg)', display)
                    if m:
                        imagenes.append({"path": Path(m.group(1)), "prompt": variante})
                else:
                    return f"Error generando imagen {i+1}: {result}"
            if not imagenes:
                return "No pude generar las imagenes."
        else:
            if not IMAGES_DIR.exists():
                return "No hay imagenes generadas todavia."
            files = sorted(IMAGES_DIR.glob("*.jpg"), key=lambda p: p.stat().st_mtime, reverse=True)[:count]
            for f in files:
                imagenes.append({"path": f, "prompt": ""})

        # Descripciones con LLM
        descripciones = []
        for img in imagenes:
            if img["prompt"]:
                desc = self._describe_image(img["prompt"])
            else:
                desc = ""
            descripciones.append(desc)

        titulo = title or (f"Imagenes: {prompt[:60]}" if prompt else "Imagenes generadas")

        office_dir = ROOT / "sandbox" / "office"
        office_dir.mkdir(parents=True, exist_ok=True)
        slug = self._slug(prompt or "imagenes")
        ts = int(datetime.now().timestamp())
        out = office_dir / f"imagenes_{slug}_{ts}.docx"

        summary = f"Crear Word con {len(imagenes)} imagen(es) incrustadas"
        if not confirmation.require("image", "to_word", summary):
            return "Cancelado."

        try:
            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "Calibri"
            style.font.size = Pt(11)

            h = doc.add_heading(titulo, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for i, img in enumerate(imagenes):
                path = img["path"]
                if not path.exists():
                    continue

                if len(imagenes) > 1:
                    doc.add_heading(f"Imagen {i+1}", level=1)

                try:
                    doc.add_picture(str(path), width=Inches(5.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    doc.add_paragraph(f"[No se pudo insertar: {e}]")
                    continue

                desc = descripciones[i] if i < len(descripciones) else ""
                if desc:
                    doc.add_paragraph(desc)

                if img["prompt"]:
                    p = doc.add_paragraph()
                    run = p.add_run(f"Prompt: {img['prompt']}")
                    run.italic = True
                    run.font.size = Pt(9)

                if i < len(imagenes) - 1:
                    doc.add_paragraph("")

            doc.save(str(out))
        except Exception as e:
            return f"Error creando Word: {e}"

        size_kb = out.stat().st_size // 1024
        print(f"[IMAGE] Word guardado: {out} ({len(imagenes)} imagenes)")

        return {
            "thought": f"Word con {len(imagenes)} imagenes HQ",
            "display": f"Word con imagenes: {out}\n({len(imagenes)} imagen(es), {size_kb} KB)",
            "voice": f"Listo. Cree un Word con {len(imagenes)} imagenes.",
        }

    def _describe_image(self, prompt):
        prompt_llm = f"""Describe en 1-2 frases cortas (maximo 40 palabras) una imagen que
tenga este prompt: "{prompt}"

Habla como si describieras la imagen resultante. Se concreto y visual.
NO digas "esta imagen muestra". Empieza directamente."""
        try:
            resp = ollama.chat(
                model=get_model("chat"),
                messages=[{"role": "user", "content": prompt_llm}],
                options={"temperature": 0.5},
                stream=False,
            )
            return resp["message"]["content"].strip()
        except Exception:
            return ""

    # ─── LIST ────────────────────────────────────────────────────────────

    def _list(self):
        if not IMAGES_DIR.exists():
            return "No hay imagenes generadas todavia."
        files = sorted(IMAGES_DIR.glob("*.jpg"), key=lambda p: p.stat().st_mtime, reverse=True)[:10]
        if not files:
            return "No hay imagenes generadas todavia."
        lines = [f"Ultimas {len(files)} imagenes:"]
        for f in files:
            lines.append(f"  - {f.name} ({f.stat().st_size // 1024} KB)")
        return {
            "thought": "",
            "display": "\n".join(lines),
            "voice": f"Tienes {len(files)} imagenes generadas.",
        }
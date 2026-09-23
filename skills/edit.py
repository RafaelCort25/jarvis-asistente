"""Skill de edicion: modifica archivos existentes con instrucciones en lenguaje natural."""
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

import ollama
from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from skills.base import Skill
from core.config_loader import CONFIG
from core import confirmation

ROOT = Path(__file__).resolve().parent.parent
UPLOADS_DIR = ROOT / "sandbox" / "uploads"

# Tipos soportados
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
TEXT_EXTS = {".txt", ".md", ".py", ".js", ".java", ".html", ".css",
             ".json", ".yaml", ".yml", ".csv", ".sh", ".ps1"}


class EditSkill(Skill):
    name = "edit"
    description = "Modifica archivos existentes (Word, Excel, texto) con instrucciones"

    def __init__(self):
        self.model = CONFIG["models"].get("default", "dolphin-directo")

    def run(self, action, params):
        if action == "modify":
            return self._modify(
                params.get("path", ""),
                params.get("instruction", ""),
                params.get("output", ""),
            )
        if action == "list_uploads":
            return self._list_uploads()
        return f"Accion desconocida en edit: {action}"

    # ─── HELPERS ─────────────────────────────────────────────────────────

    def _resolve_path(self, path_str):
        if not path_str:
            return None
        raw = path_str.strip().strip('"').strip("'")
        p = Path(raw)
        if not p.is_absolute():
            p = ROOT / p
        if p.exists():
            return p
        return None

    def _latest_in_uploads(self):
        """Devuelve el archivo mas reciente en sandbox/uploads/."""
        if not UPLOADS_DIR.exists():
            return None
        files = [f for f in UPLOADS_DIR.iterdir() if f.is_file()]
        if not files:
            return None
        return max(files, key=lambda p: p.stat().st_mtime)

    def _resolve_output(self, original, output_str):
        """Devuelve el path de salida. Si no se da, usa <nombre>_mod.<ext>."""
        if output_str:
            out = Path(output_str.strip().strip('"').strip("'"))
            if not out.is_absolute():
                out = ROOT / out
            if not out.suffix:
                out = out.with_suffix(original.suffix)
            return out
        return original.with_name(f"{original.stem}_mod{original.suffix}")

    def _ask_llm_for_replacements(self, instruction, content_preview, file_type):
        """Le pide al LLM que traduzca la instruccion en cambios concretos."""
        prompt = f"""El usuario quiere modificar un archivo de tipo {file_type}.
Su instruccion: "{instruction}"

Contenido del archivo para contexto:
---INICIO---
{content_preview[:3000]}
---FIN---

Devuelve un objeto JSON con la lista de reemplazos exactos.
Cada reemplazo debe tener una clave "old" (texto que existe literalmente en el archivo)
y una clave "new" (texto de reemplazo).

NUNCA inventes texto que no exista en el archivo.
Responde SOLO con el JSON."""

        schema = {
            "type": "object",
            "properties": {
                "replacements": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "old": {"type": "string"},
                            "new": {"type": "string"},
                        },
                        "required": ["old", "new"],
                    },
                },
            },
            "required": ["replacements"],
        }

        try:
            resp = ollama.chat(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                options={"temperature": 0.1},
                format=schema,
            )
            raw = resp["message"]["content"].strip()
            print(f"[EDIT] Respuesta del LLM: {raw[:300]}")
            return self._parse_json(raw)
        except Exception as e:
            return {"error": str(e)}

    def _parse_json(self, raw):
        """Extrae el JSON de la respuesta del LLM."""
        # Quitar ```json ... ```
        m = re.search(r'```(?:json)?\s*(\{[\s\S]*?\})\s*```', raw)
        if m:
            raw = m.group(1)
        else:
            # Buscar el primer { ... ultimo }
            start = raw.find("{")
            end = raw.rfind("}")
            if start >= 0 and end > start:
                raw = raw[start:end + 1]
        try:
            return json.loads(raw)
        except Exception as e:
            return {"error": f"JSON invalido: {e}", "raw": raw[:500]}

    # ─── PREVIEW ─────────────────────────────────────────────────────────

    def _preview_for_llm(self, path, file_type):
        """Devuelve texto de muestra del archivo para dar contexto al LLM."""
        try:
            if file_type == "docx":
                doc = Document(str(path))
                return "\n".join(p.text for p in doc.paragraphs if p.text.strip())[:3000]
            elif file_type == "xlsx":
                wb = load_workbook(str(path), read_only=True, data_only=True)
                lines = []
                for sn in wb.sheetnames[:2]:
                    ws = wb[sn]
                    lines.append(f"[Hoja: {sn}]")
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i >= 20:
                            break
                        cells = [str(c) if c is not None else "" for c in row]
                        lines.append(" | ".join(cells))
                wb.close()
                return "\n".join(lines)[:3000]
            else:
                return path.read_text(encoding="utf-8", errors="replace")[:3000]
        except Exception as e:
            return f"(no se pudo leer el archivo: {e})"

    # ─── APLICAR CAMBIOS POR FORMATO ─────────────────────────────────────

    def _apply_docx(self, path, out_path, replacements):
        doc = Document(str(path))
        total = 0
        for rep in replacements:
            old = rep.get("old", "")
            new = rep.get("new", "")
            if not old:
                continue
            # Recorrer parrafos y runs
            for para in doc.paragraphs:
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        total += 1
            # Tambien tablas
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if old in run.text:
                                    run.text = run.text.replace(old, new)
                                    total += 1
        doc.save(str(out_path))
        return total

    def _apply_xlsx(self, path, out_path, replacements):
        wb = load_workbook(str(path))
        total = 0
        for rep in replacements:
            old = rep.get("old", "")
            new = rep.get("new", "")
            if not old:
                continue
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and old in cell.value:
                            cell.value = cell.value.replace(old, new)
                            total += 1
        wb.save(str(out_path))
        return total

    def _apply_text(self, path, out_path, replacements):
        content = path.read_text(encoding="utf-8", errors="replace")
        total = 0
        for rep in replacements:
            old = rep.get("old", "")
            new = rep.get("new", "")
            if not old:
                continue
            count = content.count(old)
            if count > 0:
                content = content.replace(old, new)
                total += count
        out_path.write_text(content, encoding="utf-8")
        return total

    # ─── MODIFY ──────────────────────────────────────────────────────────

    def _modify(self, path_str, instruction, output_str):
        instruction = (instruction or "").strip()
        if not instruction:
            return "Dime que modificacion quieres hacer."

        # Resolver path: explicito o ultimo en uploads
        if path_str:
            path = self._resolve_path(path_str)
            if not path:
                return f"No encontre el archivo: {path_str}"
        else:
            path = self._latest_in_uploads()
            if not path:
                return (
                    "No hay archivos en sandbox/uploads/ para modificar. "
                    "Pon un archivo ahi o dime el path completo."
                )
            print(f"[EDIT] Usando el archivo mas reciente de uploads: {path.name}")

        ext = path.suffix.lower()
        if ext in DOCX_EXTS:
            file_type = "docx"
        elif ext in XLSX_EXTS:
            file_type = "xlsx"
        elif ext in TEXT_EXTS:
            file_type = "texto"
        else:
            return f"Formato no soportado: {ext}"

        try:
            path.relative_to(ROOT)
        except ValueError:
            return f"Ruta fuera del proyecto, bloqueado: {path}"

        out_path = self._resolve_output(path, output_str)
        try:
            out_path.relative_to(ROOT)
        except ValueError:
            return f"Ruta de salida fuera del proyecto: {out_path}"

        # 1. Contexto para el LLM
        print(f"[EDIT] Analizando {path.name} ({file_type})...")
        preview = self._preview_for_llm(path, file_type)

        # 2. Pedir cambios al LLM
        print(f"[EDIT] Interpretando instruccion con {self.model}...")
        result = self._ask_llm_for_replacements(instruction, preview, file_type)

        if "error" in result:
            return f"Error interpretando la instruccion: {result['error']}"

        replacements = result.get("replacements", [])
        if not replacements:
            return "El modelo no encontro cambios concretos para esa instruccion."

        # 3. Preview + confirmacion
        preview_lines = [f"Cambios propuestos ({len(replacements)}):"]
        for r in replacements[:10]:
            old = r.get("old", "")[:60]
            new = r.get("new", "")[:60]
            preview_lines.append(f'  "{old}" -> "{new}"')
        if len(replacements) > 10:
            preview_lines.append(f"  ... (+{len(replacements)-10} mas)")
        print("\n[EDIT] " + "\n       ".join(preview_lines) + "\n")

        summary = f"Modificar {path.name} con {len(replacements)} cambios"
        if not confirmation.require("edit", "modify", summary):
            return "Cancelado."

        # 4. Aplicar cambios
        try:
            if file_type == "docx":
                total = self._apply_docx(path, out_path, replacements)
            elif file_type == "xlsx":
                total = self._apply_xlsx(path, out_path, replacements)
            else:
                total = self._apply_text(path, out_path, replacements)
        except Exception as e:
            return f"Error aplicando cambios: {e}"

        if total == 0:
            return (
                "Ningun cambio se aplico. El texto a buscar no coincide "
                "con el contenido del archivo."
            )

        size_kb = out_path.stat().st_size // 1024
        print(f"[EDIT] Guardado: {out_path} ({total} cambios aplicados)")

        return {
            "thought": f"{total} cambios aplicados en {file_type}",
            "display": (
                f"Archivo modificado: {out_path}\n"
                f"({total} cambios aplicados, {size_kb} KB)\n\n"
                f"Origen: {path.name}"
            ),
            "voice": f"Listo. Modifique {path.name} con {total} cambios.",
        }

    # ─── LIST UPLOADS ────────────────────────────────────────────────────

    def _list_uploads(self):
        if not UPLOADS_DIR.exists():
            return "No hay carpeta de uploads todavia."

        files = [f for f in UPLOADS_DIR.iterdir() if f.is_file()]
        if not files:
            return "La carpeta sandbox/uploads/ esta vacia."

        files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        lines = [f"Archivos en uploads ({len(files)}):"]
        for f in files[:15]:
            size_kb = f.stat().st_size // 1024
            lines.append(f"  - {f.name} ({size_kb} KB)")

        return {
            "thought": "",
            "display": "\n".join(lines),
            "voice": f"Tienes {len(files)} archivos en uploads.",
        }
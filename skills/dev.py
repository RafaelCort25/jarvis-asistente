import os
import re
import shutil
import subprocess
import ollama
from pathlib import Path
from skills.base import Skill
from core.config_loader import CONFIG
from core import confirmation

# Extensiones de codigo soportadas
CODE_EXTS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".scala",
    ".html", ".css", ".scss", ".sass", ".less",
    ".json", ".yaml", ".yml", ".toml", ".xml", ".sql",
    ".sh", ".bash", ".ps1", ".bat",
    ".md", ".txt",
}

IGNORE_DIRS = {
    "venv", "node_modules", "__pycache__", ".git", "dist", "build",
    ".venv", "env", ".idea", ".vscode", "target", "Library",
    ".next", ".nuxt", "coverage",
}

MAX_FILE_CHARS = 8000
MAX_PROJECT_FILES = 15
RUN_TIMEOUT = 30

ROOT = Path(__file__).resolve().parent.parent


class DevSkill(Skill):
    name = "dev"
    description = "Revisa codigo, explica archivos, genera codigo, escribe y prueba"
    MAX_FIX_ATTEMPTS = 3

    def __init__(self):
        from core.model_config import get_model
        self.model = get_model("agent")

    def run(self, action, params):
        if action == "review_file":
            return self._review_file(params.get("path", ""))
        if action == "review_project":
            return self._review_project(params.get("path", ""))
        if action == "explain":
            return self._explain(params.get("path", ""))
        if action == "find_issues":
            return self._find_issues(params.get("path", ""))
        if action == "generate_code":
            return self._generate_code(
                params.get("description", ""),
                params.get("language", "python"),
            )
        if action == "write_file":
            return self._write_file(
                params.get("path", ""),
                params.get("content", ""),
            )
        if action == "run_file":
            return self._run_file(params.get("path", ""))
        if action == "create_and_test":
            return self._create_and_test(
                params.get("description", ""),
                params.get("language", "python"),
                params.get("path", ""),
            )
        if action == "review_to_excel":
            return self._review_to_excel(
                params.get("path", ""),
                params.get("output", ""),
            )
        if action == "review_to_word":
            return self._review_to_word(
                params.get("path", ""),
                params.get("output", ""),
            )
        return f"Accion desconocida: {action}"

    # ─── HELPERS ─────────────────────────────────────────────────────────

    def _resolve_path(self, path_str, must_exist=True):
        """Convierte string a Path. Si must_exist=False, permite rutas nuevas."""
        if not path_str:
            return None
        raw = path_str.strip().strip('"').strip("'")
        p = Path(raw)
        if not p.is_absolute():
            p = ROOT / p
        if must_exist and not p.exists():
            return None
        return p

    def _read_file(self, path):
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                content = path.read_text(encoding="latin-1", errors="ignore")
            except Exception as e:
                return None, f"Error leyendo archivo: {e}"
        except Exception as e:
            return None, f"Error leyendo archivo: {e}"
        return content, None

    def _ask_llm(self, prompt, system=None):
        messages = []
        if system:
            messages.append({"role": "system", "content": system})
        messages.append({"role": "user", "content": prompt})
        try:
            response = ollama.chat(
                model=self.model,
                messages=messages,
                options={"temperature": 0.2},
                stream=False,
            )
            return response["message"]["content"].strip()
        except Exception as e:
            return f"[ERROR LLM] {e}"

    def _clean_code_block(self, text):
        """Quita ```language ... ``` que el LLM suele anadir."""
        m = re.search(r"```[a-zA-Z0-9_+-]*\n?(.*?)```", text, re.DOTALL)
        if m:
            return m.group(1).strip()
        return text.replace("```", "").strip()

    def _es_dentro_del_proyecto(self, path):
        """Devuelve True si el path esta dentro de la raiz del proyecto."""
        try:
            path.relative_to(ROOT)
            return True
        except ValueError:
            return False

    # ─── REVIEW FILE ─────────────────────────────────────────────────────

    def _review_file(self, path_str):
        path = self._resolve_path(path_str)
        if not path:
            return f"No encontre el archivo: {path_str}"
        if not path.is_file():
            return f"No es un archivo: {path}"

        content, err = self._read_file(path)
        if err:
            return err

        truncated = False
        if len(content) > MAX_FILE_CHARS:
            content = content[:MAX_FILE_CHARS]
            truncated = True

        lang = path.suffix.lstrip(".")
        prompt = f"""Analiza este codigo {lang} y da un resumen ejecutivo en espanol.

ARCHIVO: {path.name}
{content}

Responde con este formato exacto:
1. **Proposito**: 1-2 frases sobre que hace el archivo.
2. **Estructura**: lista de funciones/clases principales (max 5).
3. **Bugs/Problemas**: errores o code smells que veas.
4. **Mejoras sugeridas**: 2-3 ideas concretas.

Se conciso. Si no ves bugs, di "Sin bugs evidentes"."""

        result = self._ask_llm(prompt, system="Eres un revisor de codigo senior. Hablas espanol, eres directo y practico.")
        note = f"\n\n_(Archivo truncado a {MAX_FILE_CHARS} caracteres)_" if truncated else ""

        return {
            "thought": f"Revisando {path.name} con {self.model}",
            "display": f"Revision de {path.name}:\n\n{result}{note}",
            "voice": f"Revise {path.name}. {result[:200]}",
        }

    # ─── REVIEW PROJECT ──────────────────────────────────────────────────

    def _review_project(self, path_str):
        path = self._resolve_path(path_str)
        if not path:
            return f"No encontre la carpeta: {path_str}"
        if not path.is_dir():
            return f"No es una carpeta: {path}"

        files = []
        for root, dirs, filenames in os.walk(path):
            dirs[:] = [d for d in dirs if d not in IGNORE_DIRS and not d.startswith(".")]
            for f in filenames:
                ext = Path(f).suffix.lower()
                if ext in CODE_EXTS:
                    full = Path(root) / f
                    try:
                        size = full.stat().st_size
                        files.append((full, size))
                    except Exception:
                        pass
            if len(files) >= MAX_PROJECT_FILES * 2:
                break

        if not files:
            return f"No encontre archivos de codigo en {path}"

        files.sort(key=lambda x: x[1], reverse=True)
        top_files = files[:MAX_PROJECT_FILES]

        summary_lines = [
            f"Proyecto: {path.name}",
            f"Total archivos de codigo: {len(files)}",
            "",
            "Archivos principales:",
        ]
        for f, size in top_files:
            rel = f.relative_to(path)
            summary_lines.append(f"  - {rel} ({size} bytes)")

        prompt = f"""Resume este proyecto de software en espanol.

{chr(10).join(summary_lines)}

Sin ver el contenido de cada archivo, dame:
1. **Tipo de proyecto**: que tipo de aplicacion parece ser.
2. **Stack probable**: lenguajes y frameworks.
3. **Estructura**: como esta organizado.
4. **Sugerencias**: que revisar o mejorar.

Maximo 200 palabras."""

        result = self._ask_llm(prompt, system="Eres arquitecto de software. Hablas espanol, eres conciso.")
        return {
            "thought": f"Analizando proyecto {path.name} ({len(files)} archivos)",
            "display": f"Revision de proyecto: {path.name}\n\n{result}",
            "voice": f"Analice el proyecto {path.name}. Tiene {len(files)} archivos de codigo.",
        }

    # ─── EXPLAIN ─────────────────────────────────────────────────────────

    def _explain(self, path_str):
        path = self._resolve_path(path_str)
        if not path or not path.is_file():
            return f"No encontre el archivo: {path_str}"

        content, err = self._read_file(path)
        if err:
            return err

        if len(content) > MAX_FILE_CHARS:
            content = content[:MAX_FILE_CHARS]

        lang = path.suffix.lstrip(".")
        prompt = f"""Explica este codigo {lang} paso a paso en espanol, para alguien que aprende.

{content}

Explica:
1. Que hace en general (2 frases).
2. Cada funcion/clase principal y su proposito.
3. Como se conectan las partes.

Se claro y didactico. Maximo 300 palabras."""

        result = self._ask_llm(prompt, system="Eres un profesor de programacion. Explicas simple y claro en espanol.")
        return {
            "thought": f"Explicando {path.name}",
            "display": f"Explicacion de {path.name}:\n\n{result}",
            "voice": f"Te explico {path.name}.",
        }

    # ─── FIND ISSUES ─────────────────────────────────────────────────────

    def _find_issues(self, path_str):
        path = self._resolve_path(path_str)
        if not path or not path.is_file():
            return f"No encontre el archivo: {path_str}"

        content, err = self._read_file(path)
        if err:
            return err

        if len(content) > MAX_FILE_CHARS:
            content = content[:MAX_FILE_CHARS]

        lang = path.suffix.lstrip(".")
        prompt = f"""Analiza este codigo {lang} y encuentra TODOS los problemas.

{content}

Busca especificamente:
- Bugs logicos
- Errores de seguridad (inyeccion SQL, XSS, hardcoded secrets)
- Problemas de rendimiento
- Code smells
- Manejo de errores faltante

Formato:
[SEVERIDAD] Linea XX: Descripcion del problema
-> Sugerencia: como arreglarlo

Si no hay problemas, di "Sin problemas evidentes".
Se directo y conciso."""

        result = self._ask_llm(prompt, system="Eres auditor de seguridad y code reviewer. Encuentras bugs que otros no ven. Hablas espanol.")
        return {
            "thought": f"Buscando bugs en {path.name}",
            "display": f"Bugs encontrados en {path.name}:\n\n{result}",
            "voice": f"Encontre algunos problemas en {path.name}.",
        }

    # ─── GENERATE CODE ───────────────────────────────────────────────────

    def _generate_code_raw(self, description, language):
        """Devuelve solo el texto del codigo, sin envolver en dict."""
        prompt = f"""Genera codigo en {language} que haga lo siguiente:
{description}

Reglas:
- Devuelve UNICAMENTE el codigo, sin explicaciones.
- Incluye comentarios utiles.
- Usa buenas practicas.
- Si es funcion, incluye docstring.
- Si es un script ejecutable, incluye un bloque main que NO use input(). Usa valores de ejemplo hardcodeados para poder ejecutarlo sin interaccion.
- NUNCA uses input() ni reads interactivos."""

        raw = self._ask_llm(prompt, system="Eres un programador experto. Generas codigo limpio y funcional.")
        return self._clean_code_block(raw)

    def _generate_code(self, description, language):
        if not description:
            return "No me dijiste que generar."

        codigo = self._generate_code_raw(description, language)
        return {
            "thought": f"Generando {language}",
            "display": f"Codigo generado:\n\n```{language}\n{codigo}\n```",
            "voice": "Listo, aqui esta el codigo.",
        }

    def _fix_code(self, codigo_actual, error_output, description, language):
        """Pide al LLM que corrija el codigo basandose en el error."""
        prompt = f"""El siguiente codigo {language} tiene un error al ejecutarse.

DESCRIPCION ORIGINAL:
{description}

CODIGO ACTUAL:
{codigo_actual}

ERROR OBTENIDO AL EJECUTAR:
{error_output}

Corrige el codigo para que funcione. Reglas:
- Devuelve UNICAMENTE el codigo corregido, sin explicaciones.
- Manten la funcionalidad original descrita.
- No uses input() ni lecturas interactivas.
- Si el error fue de sintaxis, corrige la linea exacta.
- Si el error fue logico (variables, tipos), ajusta la logica.
- Si el error fue de importacion, quita o reemplaza el import problematico."""

        raw = self._ask_llm(prompt, system="Eres un programador experto. Corriges bugs en codigo.")
        return self._clean_code_block(raw)

    # ─── WRITE FILE (con confirmacion) ───────────────────────────────────

    def _write_file(self, path_str, content):
        path = self._resolve_path(path_str, must_exist=False)
        if not path:
            return f"Ruta invalida: {path_str}"
        if not content or not content.strip():
            return "No hay contenido para escribir."

        if not self._es_dentro_del_proyecto(path):
            return f"Ruta fuera del proyecto, bloqueado por seguridad: {path}"

        existe = path.exists()
        accion = "sobrescribir" if existe else "crear"
        aviso = " (el archivo ya existe, se hara backup .bak)" if existe else ""

        summary = f"{accion} archivo: {path}{aviso}"
        if not confirmation.require("dev", "write_file", summary):
            return "Cancelado."

        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            if existe:
                backup = path.with_suffix(path.suffix + ".bak")
                shutil.copy2(path, backup)
            path.write_text(content, encoding="utf-8")
        except Exception as e:
            return f"Error escribiendo: {e}"

        return f"Archivo {'sobrescrito' if existe else 'creado'}: {path}"

    # ─── RUN FILE (con confirmacion) ─────────────────────────────────────

    def _run_file(self, path_str):
        path = self._resolve_path(path_str)
        if not path or not path.is_file():
            return f"No encontre el archivo: {path_str}"

        ext = path.suffix.lower()
        if ext == ".py":
            cmd = ["python", str(path)]
        elif ext == ".js":
            cmd = ["node", str(path)]
        elif ext == ".java":
            cmd = ["java", str(path)]
        elif ext in (".sh", ".bash"):
            cmd = ["bash", str(path)]
        elif ext == ".ps1":
            cmd = ["powershell", "-File", str(path)]
        else:
            return f"No se ejecutar {ext}. Solo .py, .js, .java, .sh, .ps1"

        summary = f"Ejecutar {path.name}: {' '.join(cmd)}"
        if not confirmation.require("dev", "run_file", summary):
            return "Cancelado."

        try:
            result = subprocess.run(
                cmd,
                cwd=str(path.parent),
                capture_output=True,
                stdin=subprocess.DEVNULL,
                text=True,
                timeout=RUN_TIMEOUT,
                encoding="utf-8",
                errors="replace",
            )
        except subprocess.TimeoutExpired:
            return f"Timeout: el archivo tardo mas de {RUN_TIMEOUT}s."
        except FileNotFoundError as e:
            return f"No encontre el interprete: {e}"
        except Exception as e:
            return f"Error ejecutando: {e}"

        stdout = (result.stdout or "").strip()
        stderr = (result.stderr or "").strip()

        def clip(s, n=1500):
            return s if len(s) <= n else s[:n] + f"\n... (recortado, {len(s)-n} mas)"

        parts = [f"$ {' '.join(cmd)}", ""]
        if stdout:
            parts.append(f"[stdout]\n{clip(stdout)}")
        if stderr:
            parts.append(f"[stderr]\n{clip(stderr)}")
        parts.append(f"[exit code: {result.returncode}]")

        return "\n".join(parts)

    # ─── HELPERS INTERNOS DEL CICLO ──────────────────────────────────────

    def _run_file_internal(self, path):
        """Ejecuta un archivo SIN pedir confirmacion (ya se pidio antes)."""
        ext = path.suffix.lower()
        if ext == ".py":
            cmd = ["python", str(path)]
        elif ext == ".js":
            cmd = ["node", str(path)]
        elif ext == ".java":
            cmd = ["java", str(path)]
        elif ext in (".sh", ".bash"):
            cmd = ["bash", str(path)]
        elif ext == ".ps1":
            cmd = ["powershell", "-File", str(path)]
        else:
            return f"(Formato {ext} no ejecutable.)"

        try:
            result = subprocess.run(
                cmd,
                cwd=str(path.parent),
                capture_output=True,
                stdin=subprocess.DEVNULL,
                text=True,
                timeout=RUN_TIMEOUT,
                encoding="utf-8",
                errors="replace",
            )
        except subprocess.TimeoutExpired:
            return f"Timeout ({RUN_TIMEOUT}s)."
        except Exception as e:
            return f"Error ejecutando: {e}"

        stdout = (result.stdout or "").strip()
        stderr = (result.stderr or "").strip()
        parts = []
        if stdout:
            parts.append(stdout[:1200])
        if stderr:
            parts.append(f"[stderr]\n{stderr[:800]}")
        parts.append(f"[exit code: {result.returncode}]")
        return "\n".join(parts)

    def _preview_code(self, codigo, iteracion=1):
        """Muestra el codigo propuesto por consola."""
        preview = codigo if len(codigo) <= 1500 else codigo[:1500] + "\n... (recortado)"
        print(f"\n[DEV] Codigo propuesto (intento {iteracion}, {len(codigo)} chars):\n")
        print(preview)
        print()

    def _write_code_to_disk(self, path, codigo):
        """Escribe el codigo en disco, con backup .bak si existe."""
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            if path.exists():
                backup = path.with_suffix(path.suffix + ".bak")
                shutil.copy2(path, backup)
            path.write_text(codigo, encoding="utf-8")
            return True
        except Exception as e:
            print(f"[DEV] Error escribiendo: {e}")
            return False

    # ─── CREATE AND TEST (ciclo completo con correccion) ─────────────────

    def _create_and_test(self, description, language, path_str):
        if not description:
            return "No me dijiste que crear."
        if not path_str:
            return "Necesito un path donde guardar el archivo."

        path = self._resolve_path(path_str, must_exist=False)
        if not path:
            return f"Ruta invalida: {path_str}"

        if not self._es_dentro_del_proyecto(path):
            return f"Ruta fuera del proyecto, bloqueado por seguridad: {path}"

        ext = path.suffix.lower()
        ejecutable = ext in (".py", ".js", ".java", ".sh", ".ps1")

        # 1. Generar codigo inicial
        print(f"[DEV] Generando {language} con {self.model}...")
        codigo = self._generate_code_raw(description, language)
        if not codigo or codigo.startswith("[ERROR LLM]"):
            return f"Error generando codigo: {codigo}"

        self._preview_code(codigo, iteracion=1)
        summary = f"Escribir {language} en {path} ({len(codigo)} chars)"
        if not confirmation.require("dev", "create_and_test", summary):
            return "Cancelado por el usuario."

        if not self._write_code_to_disk(path, codigo):
            return f"Error escribiendo {path}"

        print(f"[DEV] Archivo escrito: {path}")

        if not ejecutable:
            return f"Codigo guardado en {path}. (Formato no ejecutable, no se corre test.)"

        # 2. Ciclo: test + correccion
        for intento in range(1, self.MAX_FIX_ATTEMPTS + 1):
            if not confirmation.require(
                "dev", "run_file", f"Ejecutar {path.name} (intento {intento})"
            ):
                return f"Codigo guardado en {path}. Test cancelado en intento {intento}."

            resultado = self._run_file_internal(path)
            print(f"[DEV] Intento {intento} resultado:\n{resultado}\n")

            if "[exit code: 0]" in resultado:
                if intento == 1:
                    return f"Codigo guardado en {path}.\n\nTest exitoso:\n{resultado}"
                return (
                    f"Codigo guardado en {path}.\n\n"
                    f"Test exitoso tras {intento} intentos:\n{resultado}"
                )

            if intento == self.MAX_FIX_ATTEMPTS:
                return (
                    f"Codigo guardado en {path}, pero sigue fallando tras "
                    f"{self.MAX_FIX_ATTEMPTS} intentos.\n\nUltimo resultado:\n{resultado}"
                )

            print("[DEV] Error detectado. Pidiendo correccion al modelo...")
            codigo_corregido = self._fix_code(codigo, resultado, description, language)
            if not codigo_corregido or codigo_corregido.startswith("[ERROR LLM]"):
                return f"Error pidiendo correccion: {codigo_corregido}"

            if codigo_corregido.strip() == codigo.strip():
                return (
                    f"El modelo no encontro cambios que hacer. "
                    f"El test sigue fallando:\n{resultado}"
                )

            self._preview_code(codigo_corregido, iteracion=intento + 1)
            summary = f"Reescribir {path} con correccion {intento + 1}/{self.MAX_FIX_ATTEMPTS}"
            if not confirmation.require("dev", "write_file", summary):
                return f"Correccion rechazada. Se queda el codigo del intento {intento}."

            if not self._write_code_to_disk(path, codigo_corregido):
                return f"Error reescribiendo {path}"

            codigo = codigo_corregido

        return "Ciclo terminado sin exito."

    # ─── COMBO: REVISAR → EXCEL ──────────────────────────────────────────

    def _review_to_excel(self, path_str, output_str):
        """
        Combo: revisa un archivo, extrae bugs reales, y los escribe en un Excel.
        """
        import json as _json
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

                # Si no hay path explicito, usar el ultimo archivo de uploads
        if not path_str:
            uploads_dir = ROOT / "sandbox" / "uploads"
            if uploads_dir.exists():
                candidatos = [
                    f for f in uploads_dir.iterdir()
                    if f.is_file() and f.suffix.lower() in (".py", ".js", ".java", ".txt", ".md", ".ts", ".go", ".rb")
                ]
                if candidatos:
                    path = max(candidatos, key=lambda p: p.stat().st_mtime)
                    print(f"[DEV] Usando el ultimo archivo de uploads: {path.name}")
                else:
                    return "No hay archivos de codigo en uploads/ para revisar."
            else:
                return "No hay carpeta uploads/ y no diste un path."
        else:
            path = self._resolve_path(path_str)
            if not path or not path.is_file():
                return f"No encontre el archivo: {path_str}"

        # 1. Leer el archivo de verdad
        content, err = self._read_file(path)
        if err:
            return err
        if len(content) > MAX_FILE_CHARS:
            content = content[:MAX_FILE_CHARS]

        # 2. Pedir al LLM que encuentre bugs en JSON estructurado
        print(f"[DEV] Analizando {path.name} con {self.model}...")
        prompt = f"""Analiza el siguiente codigo y encuentra bugs o problemas REALES.

ARCHIVO: {path.name}
{content}

Devuelve un JSON con esta estructura exacta:
{{
  "bugs": [
    {{
      "linea": "numero o rango de linea (ej: 37 o 37-42) o 'N/A' si no aplica",
      "severidad": "alta | media | baja",
      "descripcion": "que problema es, concreto",
      "sugerencia": "como arreglarlo, concreto"
    }}
  ]
}}

Reglas:
- SOLO bugs reales que veas en el codigo. NO inventes.
- Si no hay bugs, devuelve {{"bugs": []}}.
- Severidad: "alta" para bugs graves (crash, seguridad), "media" para problemas logicos, "baja" para estilo.
- Maximo 15 bugs.
- Ordena de mayor a menor severidad.
- NO uses markdown, solo el JSON."""

        schema = {
            "type": "object",
            "properties": {
                "bugs": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "linea": {"type": "string"},
                            "severidad": {"type": "string"},
                            "descripcion": {"type": "string"},
                            "sugerencia": {"type": "string"},
                        },
                        "required": ["linea", "severidad", "descripcion", "sugerencia"],
                    },
                },
            },
            "required": ["bugs"],
        }

        try:
            resp = ollama.chat(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                options={"temperature": 0.1},
                format=schema,
                stream=False,
            )
            raw = resp["message"]["content"].strip()
            print(f"[DEV] Bugs encontrados (raw): {raw[:300]}")
        except Exception as e:
            return f"Error del LLM analizando el codigo: {e}"

        # 3. Parsear
        try:
            data = _json.loads(raw)
        except Exception as e:
            return f"El LLM no devolvio JSON valido: {e}"

        bugs = data.get("bugs", [])
        sin_bugs = len(bugs) == 0

        # 4. Preview
        print(f"\n[DEV] Bugs a exportar ({len(bugs)}):")
        for b in bugs[:5]:
            print(f"  [{b['severidad'].upper()}] L{b['linea']}: {b['descripcion'][:60]}")
        if len(bugs) > 5:
            print(f"  ... (+{len(bugs)-5} mas)")
        print()

        # 5. Resolver path de salida
        if output_str:
            out = Path(output_str.strip().strip('"').strip("'"))
            if not out.is_absolute():
                out = ROOT / out
            if out.suffix.lower() != ".xlsx":
                out = out.with_suffix(".xlsx")
        else:
            office_dir = ROOT / "sandbox" / "office"
            office_dir.mkdir(parents=True, exist_ok=True)
            out = office_dir / f"bugs_{path.stem}.xlsx"

        # 6. Confirmacion
        summary = f"Crear Excel con {len(bugs)} bugs reales de {path.name}"
        if not confirmation.require("dev", "review_to_excel", summary):
            return "Cancelado."

        # 7. Escribir Excel
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Bugs"

            header = ["linea", "severidad", "descripcion", "sugerencia"]
            ws.append(header)
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            if sin_bugs:
                ws.append(["N/A", "info", "Sin bugs evidentes en el archivo", "El codigo parece estar bien"])
            else:
                colors = {
                    "alta": "FFCCCC",
                    "media": "FFE5B4",
                    "baja": "CCFFCC",
                }
                for b in bugs:
                    ws.append([
                        b.get("linea", "N/A"),
                        b.get("severidad", "media"),
                        b.get("descripcion", ""),
                        b.get("sugerencia", ""),
                    ])
                    row_idx = ws.max_row
                    sev = (b.get("severidad") or "media").lower()
                    fill_color = colors.get(sev, "FFFFFF")
                    for col in range(1, 5):
                        ws.cell(row=row_idx, column=col).fill = PatternFill(
                            start_color=fill_color,
                            end_color=fill_color,
                            fill_type="solid",
                        )

            widths = [10, 12, 60, 60]
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[chr(64 + i)].width = w

            out.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(out))
        except Exception as e:
            return f"Error creando Excel: {e}"

        size_kb = out.stat().st_size // 1024
        print(f"[DEV] Excel guardado: {out}")

        if sin_bugs:
            return {
                "thought": "Sin bugs detectados, Excel informativo creado",
                "display": f"Excel creado: {out}\n(Sin bugs evidentes en {path.name}, {size_kb} KB)",
                "voice": f"Revise {path.name}, sin bugs evidentes.",
            }

        return {
            "thought": f"{len(bugs)} bugs reales en {path.name}",
            "display": (
                f"Excel con bugs de {path.name}: {out}\n"
                f"({len(bugs)} bugs encontrados, {size_kb} KB)\n\n"
                f"Alta: {sum(1 for b in bugs if b['severidad'].lower() == 'alta')}, "
                f"Media: {sum(1 for b in bugs if b['severidad'].lower() == 'media')}, "
                f"Baja: {sum(1 for b in bugs if b['severidad'].lower() == 'baja')}"
            ),
            "voice": f"Listo. Encontre {len(bugs)} bugs reales en {path.name}.",
        }

    # ─── COMBO: REVISAR → WORD ───────────────────────────────────────────

    def _review_to_word(self, path_str, output_str):
        """
        Combo: revisa un archivo y genera un Word con el analisis completo.
        """
        from docx import Document
        from docx.shared import Pt

        if not path_str:
            uploads_dir = ROOT / "sandbox" / "uploads"
            if uploads_dir.exists():
                candidatos = [
                    f for f in uploads_dir.iterdir()
                    if f.is_file() and f.suffix.lower() in (".py", ".js", ".java", ".txt", ".md", ".ts", ".go", ".rb")
                ]
                if candidatos:
                    path = max(candidatos, key=lambda p: p.stat().st_mtime)
                    print(f"[DEV] Usando el ultimo archivo de uploads: {path.name}")
                else:
                    return "No hay archivos de codigo en uploads/ para revisar."
            else:
                return "No hay carpeta uploads/ y no diste un path."
        else:
            path = self._resolve_path(path_str)
            if not path or not path.is_file():
                return f"No encontre el archivo: {path_str}"

        content, err = self._read_file(path)
        if err:
            return err
        if len(content) > MAX_FILE_CHARS:
            content = content[:MAX_FILE_CHARS]

        print(f"[DEV] Analizando {path.name} con {self.model}...")
        prompt = f"""Analiza el siguiente codigo y escribe un informe en espanol.

ARCHIVO: {path.name}
{content}

Estructura del informe:
# Analisis de {path.name}

## Proposito
1-2 frases sobre que hace el archivo.

## Estructura
Lista de funciones y clases principales (max 8).

## Bugs y problemas
Lista de bugs REALES que veas. Si no hay, di "Sin bugs evidentes".

## Mejoras sugeridas
2-4 sugerencias concretas.

Se directo y tecnico. Maximo 500 palabras."""

        try:
            resp = ollama.chat(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                options={"temperature": 0.3},
                stream=False,
            )
            analisis = resp["message"]["content"].strip()
        except Exception as e:
            return f"Error del LLM: {e}"

        if output_str:
            out = Path(output_str.strip().strip('"').strip("'"))
            if not out.is_absolute():
                out = ROOT / out
            if out.suffix.lower() != ".docx":
                out = out.with_suffix(".docx")
        else:
            office_dir = ROOT / "sandbox" / "office"
            office_dir.mkdir(parents=True, exist_ok=True)
            out = office_dir / f"analisis_{path.stem}.docx"

        summary = f"Crear Word con analisis de {path.name}"
        if not confirmation.require("dev", "review_to_word", summary):
            return "Cancelado."

        try:
            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "Calibri"
            style.font.size = Pt(11)

            for line in analisis.split("\n"):
                line = line.rstrip()
                if not line.strip():
                    continue
                if line.startswith("# "):
                    doc.add_heading(line[2:].strip(), level=0)
                elif line.startswith("## "):
                    doc.add_heading(line[3:].strip(), level=1)
                elif line.startswith("### "):
                    doc.add_heading(line[4:].strip(), level=2)
                elif line.startswith("- ") or line.startswith("* "):
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
                elif re.match(r'^\d+\.\s', line):
                    doc.add_paragraph(re.sub(r'^\d+\.\s', '', line), style="List Number")
                else:
                    doc.add_paragraph(line.strip())

            out.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(out))
        except Exception as e:
            return f"Error creando Word: {e}"

        size_kb = out.stat().st_size // 1024
        print(f"[DEV] Word guardado: {out}")

        return {
            "thought": f"Analisis de {path.name} guardado en Word",
            "display": (
                f"Word con analisis: {out}\n"
                f"({size_kb} KB)\n\n"
                f"Archivo original: {path.name}"
            ),
            "voice": f"Listo. Guarde el analisis de {path.name} en un Word.",
        }
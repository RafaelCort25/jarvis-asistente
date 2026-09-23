from pathlib import Path
from skills.base import Skill
from core.rag import RAG
from core import confirmation

ROOT = Path(__file__).resolve().parent.parent

class DocsSkill(Skill):
    name = "docs"
    description = "Indexa y consulta documentos (PDF, TXT, MD, DOCX)"

    def __init__(self):
        self.rag = RAG()

    def run(self, action, params):
        if action == "ask_to_word":
            return self._ask_to_word(
                params.get("query", ""),
                params.get("title", ""),
            )
        if action == "index_file":
            return self.rag.index_file(params.get("path", ""))
        if action == "index_folder":
            return self.rag.index_folder(params.get("path", ""))
        if action == "ask":
            result = self.rag.ask(params.get("query", ""))
            answer = result["answer"]
            sources = result["sources"]
            if sources:
                return f"{answer}\n\nFuentes: {', '.join(sources)}"
            return answer
        if action == "list":
            return self.rag.list_documents()
        if action == "delete":
            return self.rag.delete_document(params.get("name", ""))
        return f"Accion desconocida en docs: {action}"
    def _ask_to_word(self, query, title):
        """Combo: consulta el RAG y guarda la respuesta en un Word."""
        from docx import Document
        from docx.shared import Pt

        query = (query or "").strip()
        if not query:
            return "Dime sobre que quieres que consulte el RAG."

        # 1. Consultar el RAG (reutilizando el metodo existente)
        print(f"[DOCS] Consultando RAG: {query[:60]}...")
        result = self.rag.ask(query)

        answer = result.get("answer", "")
        sources = result.get("sources", [])

        if not answer:
            return "El RAG no devolvio respuesta."

        # 2. Resolver titulo y path
        titulo = title or f"Consulta: {query[:60]}"
        office_dir = ROOT / "sandbox" / "office"
        office_dir.mkdir(parents=True, exist_ok=True)

        import re as _re
        slug = _re.sub(r'[^a-z0-9]+', '_', query.lower())[:50].strip("_")
        from datetime import datetime as _dt
        ts = int(_dt.now().timestamp())
        out = office_dir / f"rag_{slug}_{ts}.docx"

        from core import confirmation
        summary = f"Crear Word con respuesta del RAG ({len(answer)} chars)"
        if not confirmation.require("docs", "ask_to_word", summary):
            return "Cancelado."

        # 3. Crear el Word
        try:
            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "Calibri"
            style.font.size = Pt(11)

            doc.add_heading(titulo, level=0)

            # Pregunta original
            p = doc.add_paragraph()
            run = p.add_run(f"Pregunta: {query}")
            run.italic = True

            doc.add_paragraph("")

            # Respuesta del RAG
            doc.add_heading("Respuesta", level=1)
            for line in answer.split("\n"):
                line = line.rstrip()
                if not line.strip():
                    continue
                if line.startswith("- ") or line.startswith("* "):
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
                elif _re.match(r'^\d+\.\s', line):
                    doc.add_paragraph(_re.sub(r'^\d+\.\s', '', line), style="List Number")
                else:
                    doc.add_paragraph(line.strip())

            # Fuentes
            if sources:
                doc.add_paragraph("")
                doc.add_heading("Fuentes consultadas", level=1)
                for s in sources:
                    doc.add_paragraph(s, style="List Bullet")

            out.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(out))
        except Exception as e:
            return f"Error creando Word: {e}"

        size_kb = out.stat().st_size // 1024
        print(f"[DOCS] Word guardado: {out}")

        return {
            "thought": f"Word creado con respuesta del RAG ({len(sources)} fuentes)",
            "display": (
                f"Word con respuesta del RAG: {out}\n"
                f"({size_kb} KB, {len(sources)} fuente(s))\n\n"
                f"Pregunta: {query}"
            ),
            "voice": f"Listo. Cree un Word con la respuesta.",
        }